from copy import deepcopy
import json
from pathlib import Path
import re
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from nexgent.organization import OrganizationService
from nexgent.organization_tools import WorkspaceTools, artifact_matches
from nexgent.organization_spreadsheets import calculate
from test_organization import Model


def tables(quantity=3):
    return [{'name': 'Summary', 'rows': [['Item', 'Total'], ['A', '=Inputs!B2*Inputs!C2']]},
            {'name': 'Inputs', 'rows': [['Item', 'Quantity', 'Price'], ['A', quantity, 10]]}]


def test_preloaded_math_accepts_normal_import_without_granting_other_imports(tmp_path):
    tools = WorkspaceTools(tmp_path, tmp_path / 'output')
    for code in ['import math\nreturn math.ceil(payload)', 'import math as m\nreturn m.ceil(payload)']:
        assert tools.run_python(code, 3.2)['value'] == 4
    with pytest.raises(ValueError):
        tools.run_python('import os\nreturn os.getcwd()', None)


def test_excel_deliverable_has_actual_calculated_caches_and_live_formulas(tmp_path):
    tools = WorkspaceTools(tmp_path, tmp_path / '.nexgent' / 'outputs' / 'writer')
    artifact = tools.call('write_spreadsheet', {'name': 'result.xlsx', 'sheets': tables()})
    assert artifact_matches(artifact)
    reader = WorkspaceTools(tmp_path, tmp_path / 'reader', shared_artifacts=[artifact], allow_artifact_writes=False)
    assert reader.call('read_spreadsheet', {'path': artifact['path'], 'sheet': 'Summary', 'cell_range': 'A1:B2'})['values'] == [['Item', 'Total'], ['A', 30]]
    assert reader.call('query_spreadsheet', {'path': artifact['path'], 'sheet': 'Summary', 'sql': 'SELECT Total FROM data'})['rows'] == [('30.0',)]
    assert reader.query_spreadsheet(artifact['path'], 'Summary', 'SELECT Total FROM Summary')['rows'] == [('30.0',)]
    with pytest.raises(ValueError, match=r'SELECT \* FROM data'):
        reader.query_spreadsheet(artifact['path'], 'Summary', 'SELECT * FROM wrong_sheet')
    book = load_workbook(artifact['path'], data_only=True)
    assert book['Summary']['B2'].value == 30
    book.close()
    book = load_workbook(artifact['path'])
    assert book['Summary']['B2'].value == '=Inputs!B2*Inputs!C2'
    book['Inputs']['B2'] = 5
    book.save(artifact['path'])
    book.close()
    assert not artifact_matches(artifact)
    # A changed Excel input invalidates the cache; the reader actually recalculates.
    assert reader.call('read_spreadsheet', {'path': artifact['path'], 'sheet': 'Summary', 'cell_range': 'B2:B2'})['values'] == [[50]]
    assert calculate(tables(0))['Summary', 2, 2] == 0
    assert 'write_spreadsheet' not in {t['name'] for t in reader.registry.describe()}


@pytest.mark.parametrize('formula', ['=1/0', '=Missing!A1', '=B2', "='[outside.xlsx]Data'!A1", '=UNKNOWNFUNCTION(A1)'])
def test_invalid_formulas_do_not_create_plausible_workbooks(tmp_path, formula):
    tools = WorkspaceTools(tmp_path, tmp_path / 'output')
    with pytest.raises(Exception):
        tools.call('write_spreadsheet', {'name': 'bad.xlsx', 'sheets': [{'name': 'Data', 'rows': [['Input', 'Result'], [2, formula]]}]})
    assert not (tmp_path / 'output' / 'bad.xlsx').exists()


def test_explicit_attachments_are_snapshots_with_exact_file_access(tmp_path):
    outside = tmp_path / 'outside'
    outside.mkdir()
    selected = outside / 'facts.csv'
    selected.write_text('Item,Amount\nA,7\n', encoding='utf-8')
    neighbor = outside / 'private.txt'
    neighbor.write_text('not selected', encoding='utf-8')
    service = OrganizationService(tmp_path / 'project', gateway_factory=Model)
    attachment = service.attach_files([selected])[0]
    selected.write_text('changed original', encoding='utf-8')
    assert Path(attachment['path']).read_text(encoding='utf-8') == 'Item,Amount\nA,7\n'
    tools = WorkspaceTools(service.root, service.root / 'out', shared_artifacts=[attachment])
    assert tools.call('query_csv', {'path': attachment['path'], 'sql': 'SELECT Amount FROM data'})['rows'] == [('7',)]
    with pytest.raises(ValueError):
        tools.source(neighbor)
    with pytest.raises(ValueError):
        service.run('bad snapshot', inputs={'attachments': [{'path': str(neighbor)}]})
    Path(attachment['path']).write_text('changed snapshot', encoding='utf-8')
    with pytest.raises(ValueError, match='changed'):
        service.run('bad snapshot', inputs={'attachments': [attachment]})


def test_excel_without_saved_dimensions_is_readable(tmp_path):
    tools = WorkspaceTools(tmp_path, tmp_path / 'outputs')
    artifact = tools.write_spreadsheet('source.xlsx', tables())
    path = Path(artifact['path'])
    with ZipFile(path) as archive:
        entries = {name: archive.read(name) for name in archive.namelist()}
    with ZipFile(path, 'w') as archive:
        for name, content in entries.items():
            if name.startswith('xl/worksheets/sheet'):
                content = re.sub(rb'<dimension\b[^>]*/>', b'', content)
            archive.writestr(name, content)
    assert tools.read_spreadsheet(str(path), 'Summary', 'B2:B2')['values'] == [[30]]


class ExcelModel(Model):
    def ask(self, role, prompt, payload, **kwargs):
        # Keep the normal recording model; its legacy computation stub knows CSV/Python.
        result = super().ask(role, prompt, {**payload, 'verification_required': False}, **kwargs)
        if 'acceptance rubric' in prompt:
            return {'criteria': 'Correct editable Excel file', 'required_artifacts': ['result.xlsx']}
        if 'Assign the task' in prompt:
            return {'recruits': [], 'peer_review': False, 'assignments': [
                {'member': 'analyst', 'task': 'Read attached source and produce Excel', 'depends_on': [], 'required_tools': ['read_spreadsheet', 'query_spreadsheet', 'write_spreadsheet']},
                {'member': 'checker', 'task': 'Read output and independently query source', 'depends_on': ['analyst'], 'required_tools': ['read_spreadsheet', 'query_spreadsheet']}]}
        if 'Perform your assigned' in prompt:
            source = payload['task']['inputs']['attachments'][0]['path']
            trace = payload['tool_results']
            assert all('error' not in t for t in trace)
            analyst = payload['assignment'] == 'Read attached source and produce Excel'
            if not trace:
                output = source if analyst else payload['upstream_results'][0]['artifacts'][0]['path']
                return {'tool': 'read_spreadsheet', 'arguments': {'path': output, 'sheet': '', 'cell_range': ''}}
            if len(trace) == 1:
                return {'tool': 'query_spreadsheet', 'arguments': {'path': source, 'sheet': 'Inputs', 'sql': 'SELECT SUM(CAST(Quantity AS REAL)*CAST(Price AS REAL)) AS total FROM data'}}
            if analyst and len(trace) == 2:
                return {'tool': 'write_spreadsheet', 'arguments': {'name': 'result.xlsx', 'sheets': tables()}}
            return {'answer': 'Calculated total is 30; output has input-linked formulas.'}
        if 'Independently evaluate' in prompt:
            assert all(t['name'] != 'write_spreadsheet' for t in payload['tools'])
            checks = payload.get('verification_results', [])
            if not checks:
                return {'tool': 'read_spreadsheet', 'arguments': {'path': payload['artifacts'][0]['path'], 'sheet': 'Summary', 'cell_range': 'A1:B2'}}
            if len(checks) == 1:
                assert checks[0]['result']['values'][1][1] == 30
                return {'tool': 'query_spreadsheet', 'arguments': {'path': payload['task']['inputs']['attachments'][0]['path'], 'sheet': 'Inputs', 'sql': 'SELECT SUM(CAST(Quantity AS REAL)*CAST(Price AS REAL)) FROM data'}}
            assert checks[1]['result']['rows'] == [(30.0,)]
            return {**result, 'required_artifacts': ['result.xlsx']}
        if role == 'improver':
            return {'organization': None, 'reason': 'Keep current team'}
        return result


def service_and_attachment(tmp_path):
    # Product writes the fixture through its normal writer, then the user attaches it.
    outside = WorkspaceTools(tmp_path, tmp_path / 'outside')
    source = outside.write_spreadsheet('source.xlsx', tables())
    service = OrganizationService(tmp_path / 'project', gateway_factory=ExcelModel)
    return service, service.attach_files([source['path']])


def test_attached_excel_completes_evaluation_and_feedback_replay(tmp_path):
    service, attachments = service_and_attachment(tmp_path)
    first = service.run('Produce result.xlsx', inputs={'attachments': attachments})
    assert first['status'] == 'completed', first.get('error')
    assert first['assessment']['accepted']
    assert [r['tool'] for r in first['assessment']['verification_evidence']] == ['read_spreadsheet', 'query_spreadsheet']
    assert artifact_matches(first['result']['artifacts'][0])
    frozen = deepcopy(service.store.get(first['id']))
    service.feedback(first['id'], 'Keep formulas editable')
    learned = service.learn(first['id'])
    assert learned['status'] == 'completed'
    assert learned['inputs'] == first['inputs']
    assert service.store.get(first['id']) == frozen
    later = service.run('Continue attachment task', inputs={'attachments': attachments}, conversation_id=first['conversation_id'])
    assert attachments[0]['path'] in {a['path'] for a in later['available_artifacts']}


def test_evaluator_cannot_accept_binary_metadata_without_reading_workbook(tmp_path):
    class Blind(ExcelModel):
        def ask(self, role, prompt, payload, **kwargs):
            if 'Independently evaluate' in prompt:
                result = Model.ask(self, role, prompt, {**payload, 'verification_required': False}, **kwargs)
                if not payload['verification_results']:
                    return {'tool': 'query_spreadsheet', 'arguments': {'path': payload['task']['inputs']['attachments'][0]['path'], 'sheet': 'Inputs', 'sql': 'SELECT SUM(CAST(Quantity AS REAL)*CAST(Price AS REAL)) FROM data'}}
                return {**result, 'required_artifacts': ['result.xlsx']}
            return super().ask(role, prompt, payload, **kwargs)
    service, attachments = service_and_attachment(tmp_path)
    service.gateway_factory = Blind
    run = service.run('Produce result.xlsx', inputs={'attachments': attachments})
    assert run['status'] == 'needs_revision'
    assert not run['assessment']['accepted']
    assert 'actual delivered Excel' in run['assessment']['feedback']


def test_evaluator_completes_missing_read_without_rebuilding_valid_delivery(tmp_path):
    class FinishRead(ExcelModel):
        def ask(self, role, prompt, payload, **kwargs):
            if 'Independently evaluate' in prompt:
                result = Model.ask(self, role, prompt, {**payload, 'verification_required': False}, **kwargs)
                if not payload['verification_results']:
                    return {'tool': 'query_spreadsheet', 'arguments': {'path': payload['task']['inputs']['attachments'][0]['path'], 'sheet': 'Inputs', 'sql': 'SELECT SUM(CAST(Quantity AS REAL)*CAST(Price AS REAL)) FROM data'}}
                if payload.get('next_action', '').startswith('Read the actual delivered workbook'):
                    if len(payload['verification_results']) == 1:
                        assert payload['remaining_verification_calls'] == 3
                        return {'tool': 'read_spreadsheet', 'arguments': {'path': payload['artifacts'][0]['path'], 'sheet': '', 'cell_range': ''}}
                return {**result, 'required_artifacts': ['result.xlsx']}
            return super().ask(role, prompt, payload, **kwargs)
    service, attachments = service_and_attachment(tmp_path)
    service.gateway_factory = FinishRead
    run = service.run('Produce result.xlsx', inputs={'attachments': attachments})
    assert run['status'] == 'completed'
    assert len(run['result']['attempts']) == 1
    assert [t['tool'] for t in run['assessment']['verification_evidence']] == ['query_spreadsheet', 'read_spreadsheet']


def test_main_file_picker_and_cli_use_same_attachment_task_flow(tmp_path, qtbot, monkeypatch, capsys):
    from nexgent.ui.organization_window import OrganizationWindow, QFileDialog
    from nexgent.cli import main
    import nexgent.organization as module
    service, attachments = service_and_attachment(tmp_path)
    monkeypatch.setattr(QFileDialog, 'getOpenFileNames', lambda *args: ([attachments[0]['path']], ''))
    window = OrganizationWindow(service.root, service)
    qtbot.addWidget(window)
    window.attach_button.click()
    assert window.pending_attachments and 'source.xlsx' in window.attachment_label.text()
    window.composer.setPlainText('Produce result.xlsx')
    window.send.click()
    assert not window.attach_button.isEnabled()
    qtbot.waitUntil(lambda: window.worker is None, timeout=15000)
    assert service.store.list()[0]['status'] == 'completed'
    assert window.attach_button.isEnabled()
    assert not window.pending_attachments
    assert 'result.xlsx' in window.messages.toPlainText()
    original = module.OrganizationService
    monkeypatch.setattr(module, 'OrganizationService', lambda root, **kwargs: original(root, gateway_factory=ExcelModel))
    assert main(['--root', str(service.root), 'run', 'Produce result.xlsx', '--organization-demo', '--attach', attachments[0]['path']]) == 0
    assert json.loads(capsys.readouterr().out)['assessment']['accepted']
