"""Small Excel adapters; parsing, calculation and export use existing libraries."""
from datetime import date, datetime
import math
import re

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.formula.tokenizer import Tokenizer


MAX_CELLS = 20000


def scalar(value):
    if hasattr(value, 'item'):
        value = value.item()
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is not None and not isinstance(value, (str, int, float, bool)):
        raise ValueError('Excel result is not a scalar value')
    if isinstance(value, float) and not math.isfinite(value):
        raise ValueError('Excel result is not finite')
    return value


def calculate(sheets):
    """Calculate bounded local formulas, without loading external workbooks."""
    import formulas
    cells, coordinates = {}, {}
    names = {s['name'].casefold() for s in sheets}
    for sheet in sheets:
        name = sheet['name'].replace("'", "''")
        for r, row in enumerate(sheet['rows'], 1):
            for c, value in enumerate(row, 1):
                key = f"'[NEXGENT.XLSX]{name}'!{get_column_letter(c)}{r}".upper()
                if isinstance(value, str) and value.startswith('='):
                    if '[' in value or ']' in value or len(value) > 2000:
                        raise ValueError('External/structured references and oversized formulas are unsupported')
                    for token in Tokenizer(value).items:
                        if token.type == 'OPERAND' and token.subtype == 'RANGE':
                            reference = token.value
                            if '!' in reference:
                                target, reference = reference.rsplit('!', 1)
                                if target.strip("'").replace("''", "'").casefold() not in names:
                                    raise ValueError('Formula references a missing worksheet')
                            bounds = range_boundaries(reference)
                            if None in bounds or bounds[2] > 100 or bounds[3] > 2000 or (bounds[2]-bounds[0]+1)*(bounds[3]-bounds[1]+1) > MAX_CELLS:
                                raise ValueError('Use bounded A1 references within 100 columns and 2000 rows; named ranges are unsupported')
                    coordinates[sheet['name'], r, c] = key
                cells[key] = value
    if not coordinates:
        return {}
    model = formulas.ExcelModel().from_dict(cells).finish(complete=False)
    solution = model.calculate()
    values = {}
    for coordinate, key in coordinates.items():
        if key not in solution:
            raise ValueError(f'Formula has missing references or a cycle: {coordinate}')
        result = solution[key].value
        if result.shape != (1, 1):
            raise ValueError('Spilled/array formulas are unsupported; use scalar helper cells')
        value = scalar(result[0, 0])
        if isinstance(value, str) and value.startswith(('#REF!', '#VALUE!', '#DIV/0!', '#NAME?', '#N/A', '#NUM!', '#NULL!')):
            raise ValueError(f'Formula error at {coordinate}: {value}')
        values[coordinate] = value
    return values


def read_book(path, *, strict=False):
    book = load_workbook(path, read_only=True, data_only=False)
    cached = load_workbook(path, read_only=True, data_only=True)
    try:
        for workbook in (book, cached):
            for sheet in workbook:
                if sheet.max_row is None or sheet.max_column is None:
                    sheet.calculate_dimension(force=True)
        if sum(s.max_row * s.max_column for s in book) > MAX_CELLS:
            raise ValueError(f'Workbook exceeds {MAX_CELLS} used cells')
        sheets = [{'name': s.title, 'rows': [[scalar(c) for c in row] for row in s.iter_rows(values_only=True)]} for s in book]
        warnings = []
        try:
            calculated = calculate(sheets)
        except Exception as exc:
            if strict:
                raise ValueError('Cannot independently calculate workbook formulas: ' + str(exc)[:300]) from exc
            calculated = None
            warnings.append('Formula values are saved Excel caches, not recalculated: ' + str(exc)[:300])
        for sheet in sheets:
            values = [row[:] for row in sheet['rows']]
            for r, row in enumerate(sheet['rows'], 1):
                for c, value in enumerate(row, 1):
                    if isinstance(value, str) and value.startswith('='):
                        values[r-1][c-1] = (calculated[sheet['name'], r, c] if calculated is not None
                                             else scalar(cached[sheet['name']].cell(r, c).value))
            sheet['values'] = values
        return sheets, warnings
    finally:
        book.close()
        cached.close()


def preview(path, sheet='', cell_range=''):
    sheets, warnings = read_book(path)
    selected = next((s for s in sheets if s['name'] == sheet), None) if sheet else sheets[0]
    if selected is None:
        raise ValueError('Worksheet does not exist; use a listed exact sheet name')
    rows = selected['rows']
    address = cell_range or f'A1:{get_column_letter(min(len(rows[0]), 20))}{min(len(rows), 50)}'
    left, top, right, bottom = range_boundaries(address)
    if None in (left, top, right, bottom) or min(left, top) < 1 or (bottom-top+1)*(right-left+1) > 2000:
        raise ValueError('Read a bounded A1 range of at most 2000 cells')
    def slice_rows(matrix):
        return [[matrix[r-1][c-1] if r <= len(matrix) and c <= len(matrix[r-1]) else None
                 for c in range(left, right+1)] for r in range(top, bottom+1)]
    return {'path': str(path), 'sheets': [{'name': s['name'], 'rows': len(s['rows']), 'columns': len(s['rows'][0])} for s in sheets],
            'sheet': selected['name'], 'range': address, 'values': slice_rows(selected['values']),
            'cells': slice_rows(rows), 'warnings': warnings,
            'truncated': bottom < len(rows) or right < len(rows[0])}


def write_book(path, sheets):
    import xlsxwriter
    if not isinstance(sheets, list) or not 1 <= len(sheets) <= 8:
        raise ValueError('Provide one to eight worksheets')
    count, names = 0, set()
    for sheet in sheets:
        name, rows = sheet.get('name'), sheet.get('rows')
        if not isinstance(name, str) or not 1 <= len(name) <= 31 or re.search(r"[\[\]:*?/\\']", name) or name.casefold() in names:
            raise ValueError('Use unique worksheet names of 1–31 characters without Excel special characters or apostrophes')
        names.add(name.casefold())
        if not isinstance(rows, list) or not rows or len(rows) > 2000 or any(not isinstance(row, list) or not 1 <= len(row) <= 100 for row in rows):
            raise ValueError('Each sheet needs nonempty rows, at most 2000 rows and 100 columns')
        count += sum(len(row) for row in rows)
        for row in rows:
            for value in row:
                if scalar(value) != value or (isinstance(value, str) and len(value) > 30000):
                    raise ValueError('Cells must be bounded JSON scalar values')
    if count > MAX_CELLS:
        raise ValueError(f'Workbook exceeds {MAX_CELLS} cells')
    values = calculate(sheets)  # Fail before creating a file; never fabricate formula caches.
    with xlsxwriter.Workbook(path, {'strings_to_urls': False}) as book:
        body = book.add_format({'font_name': 'Arial', 'font_size': 10, 'valign': 'vcenter'})
        header = book.add_format({'font_name': 'Arial', 'font_size': 10, 'bold': True,
                                 'bg_color': '#24485C', 'font_color': '#FFFFFF', 'align': 'center', 'valign': 'vcenter'})
        number = book.add_format({'font_name': 'Arial', 'font_size': 10, 'num_format': '#,##0.00', 'valign': 'vcenter'})
        book.set_calc_mode('auto')
        for sheet in sheets:
            ws = book.add_worksheet(sheet['name'])
            ws.hide_gridlines(2)
            ws.freeze_panes(1, 0)
            width = max(map(len, sheet['rows']))
            ws.set_column(0, width-1, 22, body)
            ws.set_row(0, 26)
            for r, row in enumerate(sheet['rows']):
                for c, value in enumerate(row):
                    fmt = header if r == 0 else number if isinstance(value, (float, int)) and not isinstance(value, bool) else body
                    if isinstance(value, str) and value.startswith('='):
                        ws.write_formula(r, c, value, number if r else header, values[sheet['name'], r+1, c+1])
                    else:
                        ws.write(r, c, value, fmt)
            ws.autofit(max_width=320)
    return {'sheets': [{'name': s['name'], 'rows': len(s['rows']), 'columns': max(map(len, s['rows']))} for s in sheets],
            'formula_count': len(values)}
